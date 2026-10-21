"""PDF/DOCX export helpers used by the web application.

All Qt based rendering (SVG -> PNG, HTML -> PDF) must run in one single thread,
because the QGuiApplication is bound to the thread that created it. The web app
therefore funnels every export through ``run_in_qt_thread``.
"""

import os
import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html import escape as html_escape
from pathlib import Path

import db_report

PROJECT_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = PROJECT_ROOT / "data" / "templates" / "TeamReportTemplate.docx"

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_qt_app = None


def _init_qt() -> None:
    global _qt_app
    from PySide6.QtGui import QGuiApplication

    _qt_app = QGuiApplication.instance() or QGuiApplication(["pubquizstats"])


_qt_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="qt", initializer=_init_qt)


def run_in_qt_thread(fn, *args, **kwargs):
    """Run ``fn`` in the dedicated Qt thread and return its result (or raise)."""
    return _qt_executor.submit(fn, *args, **kwargs).result()


def submit_in_qt_thread(fn, *args, **kwargs):
    return _qt_executor.submit(fn, *args, **kwargs)


def safe_filename(name: str, fallback: str = "team") -> str:
    safe = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in name.strip()).strip("_")
    return safe or fallback


def detect_office_converter() -> str | None:
    """Return absolute path to soffice/libreoffice if available."""
    for name in ("soffice", "libreoffice"):
        resolved = shutil.which(name)
        if resolved:
            return resolved

    for path in (
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
        "/snap/bin/soffice",
        "/opt/libreoffice/program/soffice",
    ):
        if Path(path).exists():
            return path

    return None


def _format_points_cell(value) -> str:
    if value is None:
        return "-"
    as_float = float(value)
    if as_float.is_integer():
        return str(int(as_float))
    return f"{as_float:.1f}"


# ---------------------------------------------------------------------------
# HTML -> PDF (Qt)
# ---------------------------------------------------------------------------

def _html_to_pdf(html: str, path: str, landscape: bool = False) -> None:
    from PySide6.QtGui import QPageLayout, QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    document = QTextDocument()
    document.setHtml(html)

    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(path)
    printer.setPageOrientation(QPageLayout.Orientation.Landscape if landscape else QPageLayout.Orientation.Portrait)
    document.print_(printer)


def text_report_pdf(text: str, path: str) -> None:
    html = (
        "<html><body><pre style=\"font-family: monospace; font-size: 9pt;\">"
        f"{html_escape(text)}</pre></body></html>"
    )
    _html_to_pdf(html, path)


def build_theme_difficulty_html(report: dict, year: int) -> str:
    rows = report.get("events", [])
    title = f"Theme Difficulties Report ({year})"

    event_blocks = []
    for index, row in enumerate(rows):
        category_lines = [
            (
                "Puzzle",
                row.get("puzzle_category") or "-",
                row.get("puzzle_solution") or "-",
                _format_points_cell(row.get("puzzle_average_points")),
            ),
            (
                "Bilderrunde",
                row.get("image_round_topic") or "-",
                "-",
                _format_points_cell(row.get("image_average_points")),
            ),
            (
                "Überraschung",
                row.get("surprise_round_topic") or "-",
                "-",
                _format_points_cell(row.get("surprise_average_points")),
            ),
        ]
        lines_html = "".join(
            f"<div class='category-line'><span class='category-name'>{html_escape(label)}:</span> "
            f"{html_escape(avg)} Pts / {html_escape(theme)}"
            f"{(' / ' + html_escape(solution)) if label == 'Puzzle' and solution and solution != '-' else ''}</div>"
            for label, theme, solution, avg in category_lines
        )
        event_blocks.append(
            f"""
            <div class="event-block" style="page-break-inside: avoid; break-inside: avoid;">
              <div class="event-header">{html_escape(str(row.get('event_date') or '-'))} {html_escape(str(row.get('location') or '-'))}</div>
              {lines_html}
            </div>
            """
        )
        if (index + 1) % 5 == 0 and (index + 1) < len(rows):
            event_blocks.append(
                '<div class="page-break-spacer" style="page-break-after: always; break-after: page; height: 0; margin: 0;">&nbsp;</div>'
            )

    return f"""
<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <style>
    @page {{ size: A4 portrait; margin: 12mm; }}
    body {{ font-family: Arial, sans-serif; color: #111827; font-size: 11pt; line-height: 1.5; }}
    h1 {{ margin: 0 0 6mm 0; font-size: 18pt; font-weight: 700; }}
    p.meta {{ margin: 0 0 6mm 0; color: #4b5563; }}
    .event-block {{ margin: 0 0 14px 0; padding: 0 0 8px 0; page-break-inside: avoid; break-inside: avoid; page-break-before: auto; }}
    .event-header {{ font-size: 19pt; font-weight: 700; margin: 0 0 8px 0; }}
    .category-line {{ margin: 0 0 6px 0; font-size: 14pt; }}
    .category-name {{ font-weight: 700; }}
  </style>
</head>
<body>
  <h1>{html_escape(title)}</h1>
  <p class="meta">Generated: {html_escape(datetime.now().strftime('%Y-%m-%d %H:%M'))} | Events: {len(rows)}</p>
  {''.join(event_blocks)}
</body>
</html>
"""


def build_team_category_detail_html(report: dict, year: int) -> str:
    rows = report.get("events", [])
    title = f"Team Category Detail Report - {report['team_name']} ({year})"

    table_rows = []
    for row in rows:
        table_rows.append(
            "<tr>"
            f"<td>{html_escape(str(row.get('event_date') or '-'))}</td>"
            f"<td>{html_escape(str(row.get('puzzle_category') or '-'))}</td>"
            f"<td>{html_escape(str(row.get('puzzle_solution') or '-'))}</td>"
            f"<td class='num'>{html_escape(_format_points_cell(row.get('puzzle_points')))}</td>"
            f"<td>{html_escape(str(row.get('image_round_category') or '-'))}</td>"
            f"<td class='num'>{html_escape(_format_points_cell(row.get('image_round_points')))}</td>"
            f"<td>{html_escape(str(row.get('surprise_round_category') or '-'))}</td>"
            f"<td class='num'>{html_escape(_format_points_cell(row.get('surprise_round_points_no_bonus')))}</td>"
            "</tr>"
        )

    return f"""
<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <style>
    @page {{ size: A4 landscape; margin: 14mm; }}
    body {{ font-family: Arial, sans-serif; color: #111827; font-size: 10pt; }}
    h1 {{ margin: 0 0 6mm 0; font-size: 17pt; }}
    p.meta {{ margin: 0 0 6mm 0; color: #4b5563; }}
    table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
    th, td {{ border: 1px solid #d1d5db; padding: 6px 7px; vertical-align: top; word-wrap: break-word; }}
    th {{ background: #f3f4f6; text-align: left; font-weight: 700; }}
    td.num {{ text-align: right; }}
    th:nth-child(1), td:nth-child(1) {{ width: 12%; }}
    th:nth-child(2), td:nth-child(2) {{ width: 12%; }}
    th:nth-child(3), td:nth-child(3) {{ width: 17%; }}
    th:nth-child(4), td:nth-child(4) {{ width: 8%; }}
    th:nth-child(5), td:nth-child(5) {{ width: 14%; }}
    th:nth-child(6), td:nth-child(6) {{ width: 8%; }}
    th:nth-child(7), td:nth-child(7) {{ width: 18%; }}
    th:nth-child(8), td:nth-child(8) {{ width: 11%; }}
  </style>
</head>
<body>
  <h1>{html_escape(title)}</h1>
  <p class="meta">Generated: {html_escape(datetime.now().strftime('%Y-%m-%d %H:%M'))} | Events: {len(rows)}</p>
  <table>
    <thead>
      <tr>
        <th>Event Date</th>
        <th>Puzzle Category</th>
        <th>Puzzle Solution</th>
        <th>Puzzle Points</th>
        <th>Image Round Category</th>
        <th>Image Points</th>
        <th>Surprise Round Category</th>
        <th>Surprise Points (no bonus)</th>
      </tr>
    </thead>
    <tbody>
      {''.join(table_rows)}
    </tbody>
  </table>
</body>
</html>
"""


def theme_difficulty_pdf(year: int, db_path: str, path: str) -> None:
    report = db_report.get_theme_difficulty_report(year=year, db_path=db_path)
    _html_to_pdf(build_theme_difficulty_html(report, year), path)


def team_category_detail_pdf(team_name: str, year: int, db_path: str, path: str) -> None:
    report = db_report.get_team_event_theme_report(team_name=team_name, db_path=db_path, year=year)
    _html_to_pdf(build_team_category_detail_html(report, year), path, landscape=True)


# ---------------------------------------------------------------------------
# Team report: DOCX template -> PDF
# ---------------------------------------------------------------------------

def _template_values(report: dict) -> dict:
    events = report.get("season_events", [])
    radar_rounds = report.get("radar_rounds", [])
    best_result = report.get("best_result")

    best_pos_event = None
    positioned_events = [event for event in events if event.get("position") is not None]
    if positioned_events:
        best_pos_event = min(positioned_events, key=lambda event: event["position"])

    ranked_rounds = [row for row in radar_rounds if row.get("position") is not None]
    best_category = None
    if ranked_rounds:
        best_category = min(
            ranked_rounds,
            key=lambda row: (
                row["position"],
                -(row["avg_points"] if row.get("avg_points") is not None else 0),
                row["round_name"],
            ),
        )

    year_events_count = report.get("year_events_count")
    if year_events_count is None:
        year_events_count = len(events)

    first_places = sum(1 for event in events if event.get("position") == 1)
    second_places = sum(1 for event in events if event.get("position") == 2)
    third_places = sum(1 for event in events if event.get("position") == 3)

    return {
        "Date": datetime.now().strftime("%d.%m.%Y"),
        "YEAR": str(report["year"]),
        "Team Name": report["team_name"],
        "Teams": str(report.get("teams_total", "-")),
        "ChampionshipPoints": "-" if report["championship_points"] is None else str(report["championship_points"]),
        "ChampionshipPosition": "-" if report["championship_place"] is None else str(report["championship_place"]),
        "BestPos": "-" if best_pos_event is None else str(best_pos_event["position"]),
        "DateBestPos": "-" if best_pos_event is None else best_pos_event["event_date"],
        "LocBestPos": "-" if best_pos_event is None else best_pos_event["location"],
        "BestPts": "-" if best_result is None else str(best_result["total_points"]),
        "DateBestPts": "-" if best_result is None else best_result["event_date"],
        "LocBestPts": "-" if best_result is None else best_result["location"],
        "AvgPoints": "-" if report["average_points"] is None else f"{report['average_points']:.2f}",
        "BonusEff": "-" if report.get("bonus_efficiency_avg") is None else f"{report['bonus_efficiency_avg'] * 100:.1f}%",
        "BestCat": "-" if best_category is None else best_category["round_name"],
        "BestCatPos": "-" if best_category is None else str(best_category["position"]),
        "nTeams": "-" if best_category is None else str(best_category["total_teams"]),
        "Participations": str(report["participation_count"]),
        "nRoundsPlayed": str(year_events_count),
        "1": str(first_places),
        "2": str(second_places),
        "3": str(third_places),
    }


def _iter_table_paragraphs(tables):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _iter_docx_paragraphs(document):
    yield from document.paragraphs
    yield from _iter_table_paragraphs(document.tables)

    # Include placeholders in headers/footers (e.g. footer date tokens).
    for section in document.sections:
        yield from section.header.paragraphs
        yield from _iter_table_paragraphs(section.header.tables)
        yield from section.footer.paragraphs
        yield from _iter_table_paragraphs(section.footer.tables)


def _locate_span(run_texts: list[str], start: int, end: int):
    boundaries = []
    offset = 0
    for idx, text in enumerate(run_texts):
        boundaries.append((idx, offset, offset + len(text)))
        offset += len(text)

    start_idx = None
    end_idx = None
    for idx, b_start, b_end in boundaries:
        if start_idx is None and b_start <= start < b_end:
            start_idx = idx
        if b_start < end <= b_end:
            end_idx = idx
            break
    return boundaries, start_idx, end_idx


def _replace_placeholder_in_paragraph(paragraph, placeholder: str, value: str) -> bool:
    run_texts = [run.text for run in paragraph.runs]
    full_text = "".join(run_texts)
    if placeholder not in full_text:
        return False

    start = full_text.find(placeholder)
    end = start + len(placeholder)
    boundaries, start_idx, end_idx = _locate_span(run_texts, start, end)

    if start_idx is None or end_idx is None:
        paragraph.text = full_text.replace(placeholder, value)
        return True

    start_run = paragraph.runs[start_idx]
    end_run = paragraph.runs[end_idx]
    prefix = start_run.text[: start - boundaries[start_idx][1]]
    suffix = end_run.text[end - boundaries[end_idx][1]:]
    start_run.text = prefix + value + suffix

    for idx in range(start_idx + 1, end_idx + 1):
        paragraph.runs[idx].text = ""
    return True


def _replace_placeholder_across_text_nodes(text_nodes, placeholder: str, value: str) -> bool:
    """Replace one placeholder occurrence across XML text nodes while preserving layout."""
    run_texts = [node.text or "" for node in text_nodes]
    full_text = "".join(run_texts)
    start = full_text.find(placeholder)
    if start < 0:
        return False
    end = start + len(placeholder)
    boundaries, start_idx, end_idx = _locate_span(run_texts, start, end)
    if start_idx is None or end_idx is None:
        return False

    start_node = text_nodes[start_idx]
    end_node = text_nodes[end_idx]
    prefix = (start_node.text or "")[: start - boundaries[start_idx][1]]
    suffix = (end_node.text or "")[end - boundaries[end_idx][1]:]
    start_node.text = prefix + value + suffix

    for idx in range(start_idx + 1, end_idx + 1):
        text_nodes[idx].text = ""
    return True


def _replace_placeholder_in_xml_text_nodes(root_element, placeholder: str, value: str) -> None:
    # WordprocessingML (w:t) handles split runs; DrawingML (a:t) covers text boxes/shapes.
    for paragraph_xpath, text_xpath in ((".//w:p", ".//w:t"), (".//a:p", ".//a:t")):
        for paragraph in root_element.xpath(paragraph_xpath):
            text_nodes = paragraph.xpath(text_xpath)
            if not text_nodes:
                continue
            while _replace_placeholder_across_text_nodes(text_nodes, placeholder, value):
                pass


def _replace_docx_placeholders(document, values: dict) -> None:
    for placeholder, value in values.items():
        token = "{" + placeholder + "}"
        replacement = str(value)
        for paragraph in _iter_docx_paragraphs(document):
            while _replace_placeholder_in_paragraph(paragraph, token, replacement):
                continue

        # Text boxes/shapes are not exposed through document.paragraphs/tables.
        _replace_placeholder_in_xml_text_nodes(document.part.element, token, replacement)
        for section in document.sections:
            _replace_placeholder_in_xml_text_nodes(section.header.part.element, token, replacement)
            _replace_placeholder_in_xml_text_nodes(section.footer.part.element, token, replacement)


def _replace_image_in_cell(cell, image_path: str) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Emu, Inches

    preserved_width = None
    preserved_height = None
    preserved_alignment = None

    # Reuse dimensions of the existing image in this cell, if present.
    for paragraph in cell.paragraphs:
        if preserved_alignment is None:
            preserved_alignment = paragraph.alignment

        for run in paragraph.runs:
            extents = run._r.xpath(".//wp:inline/wp:extent") or run._r.xpath(".//wp:anchor/wp:extent")
            if extents:
                cx = extents[0].get("cx")
                cy = extents[0].get("cy")
                if cx and cy:
                    preserved_width = Emu(int(cx))
                    preserved_height = Emu(int(cy))
                    break
        if preserved_width is not None and preserved_height is not None:
            break

    cell.text = ""
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    paragraph.alignment = preserved_alignment or WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    if preserved_width is not None and preserved_height is not None:
        run.add_picture(image_path, width=preserved_width, height=preserved_height)
    elif cell.width:
        run.add_picture(image_path, width=int(cell.width * 0.92))
    else:
        run.add_picture(image_path, width=Inches(5.8))
    return True


def _insert_charts_into_docx(document, avg_radar_png: str, pos_radar_png: str, event_chart_png: str) -> bool:
    if len(document.tables) < 2:
        return False

    table = document.tables[1]
    if len(table.rows) < 1 or len(table.columns) < 2:
        return False

    left_ok = _replace_image_in_cell(table.cell(0, 0), avg_radar_png)
    right_ok = _replace_image_in_cell(table.cell(0, 1), pos_radar_png)
    bottom_ok = False
    if len(table.rows) >= 2:
        bottom_ok = _replace_image_in_cell(table.cell(1, 0), event_chart_png)
    return left_ok and right_ok and bottom_ok


def _render_svg_png(svg_path: str, png_path: str, width: int = 1800, height: int = 1200) -> str:
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtSvg import QSvgRenderer

    renderer = QSvgRenderer(svg_path)
    if not renderer.isValid():
        raise RuntimeError("Chart SVG could not be loaded.")

    image = QImage(width, height, QImage.Format.Format_ARGB32)
    image.fill(Qt.GlobalColor.white)

    painter = QPainter(image)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing, True)
    painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform, True)
    renderer.render(painter)
    painter.end()

    if not image.save(png_path, "PNG"):
        raise RuntimeError("Failed to render chart PNG.")
    return png_path


def _convert_docx_to_pdf(docx_path: str, pdf_path: str) -> None:
    converter = detect_office_converter()
    if converter is None:
        raise RuntimeError(
            "No DOCX-to-PDF converter found. Install LibreOffice and ensure 'soffice' is available in PATH."
        )

    out_dir = Path(pdf_path).parent.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    command = [converter, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(Path(docx_path).resolve())]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(f"DOCX conversion failed: {result.stderr.strip() or result.stdout.strip()}")

    converted = out_dir / (Path(docx_path).stem + ".pdf")
    if not converted.exists():
        raise RuntimeError("DOCX conversion did not produce a PDF file.")

    target = Path(pdf_path).resolve()
    if converted.resolve() != target:
        converted.replace(target)


def _render_team_pdf_from_docx_template(report: dict, output_pdf_path: str) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("python-docx is required. Install it with: pip install python-docx") from exc

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    values = _template_values(report)

    with tempfile.TemporaryDirectory(prefix="team-report-") as tmp_dir:
        tmp_docx = Path(tmp_dir) / "team_report_filled.docx"
        tmp_avg_png = Path(tmp_dir) / "radar_plot_avg.png"
        tmp_pos_png = Path(tmp_dir) / "radar_plot_pos.png"
        tmp_event_png = Path(tmp_dir) / "event_chart.png"

        _render_svg_png(report["radar_svg_path"], str(tmp_avg_png))
        _render_svg_png(report["radar_position_svg_path"], str(tmp_pos_png))
        _render_svg_png(report["event_chart_svg_path"], str(tmp_event_png), width=2200, height=1000)

        document = Document(str(TEMPLATE_PATH))
        _replace_docx_placeholders(document, values)
        if not _insert_charts_into_docx(document, str(tmp_avg_png), str(tmp_pos_png), str(tmp_event_png)):
            raise RuntimeError("Could not place charts in second table.")

        document.save(str(tmp_docx))
        _convert_docx_to_pdf(str(tmp_docx), output_pdf_path)


def team_report_pdf(team_name: str, year: int, min_events: int, db_path: str, path: str) -> None:
    """Render the DOCX-template based team report as PDF and clean up temporary SVGs."""
    report = db_report.get_team_profile_report(team_name, year, min_events, db_path)
    try:
        _render_team_pdf_from_docx_template(report, path)
    finally:
        for key in ("radar_svg_path", "radar_position_svg_path", "event_chart_svg_path"):
            if report.get(key):
                Path(report[key]).unlink(missing_ok=True)
